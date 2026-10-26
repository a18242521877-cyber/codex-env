"""把商家后台页面自己返回的 JSON / CSV / xlsx 收成统一报表。

没有一份稳定的公开接口。这里只认「一行里同时有身份和指标」这种形状，
字段对不上就留空，不补造花费或 ROI。

金额启发式（后台改口径时可能差 100 倍，要对照页面核对）：

- CSV、xlsx 里的金额按「元」
- JSON 里的整数按「分」，小数按「元」
- JSON 字符串带小数点按「元」，纯数字按「分」

测图（创意 / 主图）不并进商品花费。只有商品、没有计划行时，按场景合成一条计划，
计划花费等于其下商品之和。已经有计划行时，不再把商品加进计划。
"""

from __future__ import annotations

import csv
import io
import json
import re
from datetime import date, datetime, timedelta
from typing import Any

from app.adapters.base import (
    CampaignInfo,
    GoodsInfo,
    ImageInfo,
    MetricDay,
    ShopProfile,
    SyncPayload,
)

_DATE_KEYS = ("statDate", "stat_date", "reportDate", "日期", "统计日期", "date", "ds")
_GOODS_ID = ("goodsId", "goods_id", "商品ID")
_GOODS_NAME = ("goodsName", "goods_name", "商品名称", "商品名")
_PLAN_ID = ("planId", "adId", "campaignId", "unitId", "计划ID", "推广计划ID")
_PLAN_NAME = ("planName", "adName", "campaignName", "计划名称", "推广计划", "计划名")
_CREATIVE_ID = ("creativeId", "imageId", "创意ID")
_CREATIVE_NAME = ("creativeName", "创意名称", "主图名称", "主图")
_IMPR = ("impression", "impressions", "showCnt", "曝光量", "曝光", "展现量", "展现")
_CLICK = ("click", "clicks", "clickCnt", "点击量", "点击")
_SPEND = ("成交花费", "推广花费", "总花费", "花费", "消耗", "spend", "cost")
_ORDER = ("成交笔数", "支付订单数", "成交订单数", "订单数", "orderNum", "orders", "orderCnt")
_GMV = ("交易额", "成交额", "成交金额", "支付金额", "交易金额", "gmv")
_CART = ("addCart", "add_to_cart", "加购数", "加购")
_BUDGET = ("dailyBudget", "dayBudget", "日限额", "日预算", "预算")
_BID = ("bid", "出价", "点击出价")
_PRICE = ("goodsPrice", "price", "拼单价", "商品价格")
_STATUS = ("status", "adStatus", "状态")
_SCENE = ("scene", "推广场景", "推广类型", "场景", "计划类型")
_STAGE = ("stage", "阶段")
_MALL_ID = ("mallId", "mall_id", "shopId")
_MALL_NAME = ("mallName", "mall_name", "shopName", "店铺名称", "店铺名")
OFFICIAL_YINGXIAO_URL = "https://yingxiao.pinduoduo.com/mains/promotionOverview"
YINGXIAO_SYNC_HINT = (
    "请先在浏览器登录推广中心，再到本应用点同步。"
    f"地址：{OFFICIAL_YINGXIAO_URL} 。"
    "请用 Edge 打开该页，自己完成登录（头像或二维码都可以），并停留在推广概况或商品报表，不要关窗口。"
)


def manual_login_tail(gate: str) -> str:
    """推广中心未登录时，都指向同一个固定页面，不代点、不等扫码。"""
    if "推广" in (gate or ""):
        return YINGXIAO_SYNC_HINT
    return (
        f"请打开{gate}，在浏览器里自己完成登录，回到推广台后再点「同步」。"
        "同步只读取已保存的登录，不会代点，也不会等待扫码。"
    )


SYDNEY_SOURCE_NOTE = (
    "昨日成交额和订单来自商家后台「商品数据」，不是推广花费。"
    "推广花费需要打开推广中心页面，自己完成登录后再点「同步」。"
)
AUTH_FAILED_NOTE = (
    "推广中心尚未登录（登录态为空），测图测款未写入。"
    + manual_login_tail("推广中心页面")
    + "这次没有写入推广花费。"
)
TOKEN_INVALID_NOTE = (
    "推广中心头像登录失败：Token不合法，测图测款未写入。不要只点头像。"
    + manual_login_tail("推广中心页面")
)
MMS_LOGIN_NOTE = "商家后台尚未登录。" + manual_login_tail("商家后台登录页")
SAVED_LOGIN_NOTE = (
    "登录已保存。请回到推广台，点「同步」读取报表。"
    "同步只读取已保存的登录，不会代点，也不会等待扫码。"
)
HOME_SPEND_NOTE = (
    "商家后台首页上的推广花费已写入。这是首页汇总，不是测款或测图。测图测款未写入。"
    + manual_login_tail("推广中心页面")
)
_HOME_SPEND_KEYS = {
    "adspend",
    "adcost",
    "adspendamt",
    "promotioncost",
    "promotionspend",
    "promospend",
    "spendamt",
    "costamt",
    "推广花费",
    "广告花费",
}
_METRIC_KEY_NORMS = {
    "spend",
    "cost",
    "impression",
    "impressions",
    "click",
    "clicks",
    "showcnt",
    "clickcnt",
    "ordernum",
    "orders",
    "ordercnt",
    "gmv",
    "花费",
    "消耗",
    "成交花费",
    "推广花费",
    "总花费",
    "曝光",
    "曝光量",
    "展现",
    "展现量",
    "点击",
    "点击量",
    "成交额",
    "交易额",
    "成交金额",
    "支付金额",
}
_GOODS_KEY_NORMS = {"goodsid", "goodsname", "商品id", "商品名称", "商品名"}
_PLAN_KEY_NORMS = {
    "planid",
    "adid",
    "campaignid",
    "unitid",
    "planname",
    "adname",
    "campaignname",
    "计划id",
    "推广计划id",
    "计划名称",
    "推广计划",
    "计划名",
}
_CREATIVE_KEY_NORMS = {
    "creativeid",
    "creativename",
    "imageid",
    "创意id",
    "创意名称",
    "主图",
    "主图名称",
}


def _key_norm(key: Any) -> str:
    return str(key).strip().lower().replace("_", "").replace(" ", "")


def _structured_hit(body: Any, id_keys: set[str]) -> bool:
    """同一行上同时有身份字段和指标字段才算报表。说明文字里的「花费」不算。"""
    found = False

    def walk(node: Any, depth: int) -> None:
        nonlocal found
        if found or depth > 8:
            return
        if isinstance(node, dict):
            keys = {_key_norm(key) for key in node.keys()}
            if (keys & id_keys) and (keys & _METRIC_KEY_NORMS):
                found = True
                return
            for value in node.values():
                if isinstance(value, (dict, list)):
                    walk(value, depth + 1)
        elif isinstance(node, list):
            for item in node[:50]:
                walk(item, depth + 1)

    walk(body, 0)
    return found


def json_has_metrics(body: Any) -> bool:
    return _structured_hit(body, _GOODS_KEY_NORMS | _PLAN_KEY_NORMS | _CREATIVE_KEY_NORMS)


def json_has_creative_metrics(body: Any) -> bool:
    return _structured_hit(body, _CREATIVE_KEY_NORMS)


def json_has_promo_goods(body: Any) -> bool:
    return _structured_hit(body, _GOODS_KEY_NORMS | _PLAN_KEY_NORMS)


def body_looks_like_mikasa(body: Any) -> bool:
    """推广接口未登录时常见：mikasaErrorCode，并且 result 是空的。"""
    found = False

    def walk(node: Any, depth: int) -> None:
        nonlocal found
        if found or depth > 6:
            return
        if isinstance(node, dict):
            if "mikasaErrorCode" in node or "mikasa_error_code" in node:
                result = node.get("result", node.get("Result", "missing"))
                if result in (None, {}, []):
                    found = True
                    return
            for value in node.values():
                if isinstance(value, (dict, list)):
                    walk(value, depth + 1)
        elif isinstance(node, list):
            for item in node[:20]:
                walk(item, depth + 1)

    walk(body, 0)
    return found


def body_token_invalid(body: Any) -> bool:
    """头像换票失败。generateAccessToken 返回 Token不合法时，推广中心并没有登录。"""
    found = False

    def walk(node: Any, depth: int) -> None:
        nonlocal found
        if found or depth > 8:
            return
        if isinstance(node, dict):
            message = node.get("errorMsg", node.get("error_msg", node.get("errorMessage")))
            if isinstance(message, str) and "Token不合法" in message:
                found = True
                return
            for value in node.values():
                if isinstance(value, (dict, list)):
                    walk(value, depth + 1)
        elif isinstance(node, list):
            for item in node[:30]:
                walk(item, depth + 1)

    walk(body, 0)
    return found


def body_auth_failed(body: Any) -> bool:
    """登录态为空、头像 Token不合法，或 mikasa 空结果。这是授权失败，不是报表里没有商品。"""
    if body_looks_like_mikasa(body) or body_token_invalid(body):
        return True
    found = False

    def walk(node: Any, depth: int) -> None:
        nonlocal found
        if found or depth > 8:
            return
        if isinstance(node, dict):
            message = node.get("errorMsg", node.get("error_msg", node.get("errorMessage")))
            if isinstance(message, str) and "登录态为空" in message:
                found = True
                return
            for value in node.values():
                if isinstance(value, (dict, list)):
                    walk(value, depth + 1)
        elif isinstance(node, list):
            for item in node[:30]:
                walk(item, depth + 1)

    walk(body, 0)
    return found


def note_says_auth_failed(note: str) -> bool:
    text = note or ""
    return any(
        mark in text
        for mark in (
            "测图测款未写入",
            "登录态为空",
            "没有写入推广花费",
            "未授权",
            "授权没有完成",
            "Token不合法",
            "请先在浏览器登录推广中心",
        )
    )


def money_to_fen(value: Any, *, source: str) -> int:
    if value is None or isinstance(value, bool):
        return 0
    if isinstance(value, str):
        text = value.replace(",", "").replace("元", "").replace("¥", "").replace("￥", "").strip()
        if text in {"", "-", "--"}:
            return 0
        try:
            number = float(text)
        except ValueError:
            return 0
        if source == "json" and "." not in text:
            return int(number)
        return int(round(number * 100))
    if isinstance(value, int):
        if source == "json":
            return value
        return int(round(value * 100))
    if isinstance(value, float):
        return int(round(value * 100))
    return 0


def _plaintext_metric(value: Any) -> bool:
    if isinstance(value, bool) or value is None or value == "":
        return False
    if isinstance(value, (int, float)):
        return True
    text = str(value).strip().replace(",", "").replace("元", "").replace("%", "")
    return bool(re.fullmatch(r"-?\d+(\.\d+)?", text))


def _sydney_item_row(item: dict, day: date, *, aggregate: bool) -> dict | None:
    amount = item.get("payOrdrAmt")
    orders = item.get("payOrdrCnt")
    page_views = item.get("gpv")
    if not any(_plaintext_metric(value) for value in (amount, orders, page_views)):
        return None
    if aggregate:
        goods_id, name = "sydney-shop", "全店商品数据"
    else:
        goods_id = item.get("goodsId", item.get("goods_id"))
        name = item.get("goodsName") or item.get("goods_name") or "商品"
        if goods_id in (None, ""):
            return None
    row: dict[str, Any] = {
        "goodsid": goods_id,
        "goodsname": str(name)[:80],
        "statdate": day.isoformat(),
        "scene": "商品数据",
        "__source": "json",
    }
    if _plaintext_metric(amount):
        row["gmv"] = amount
    if _plaintext_metric(orders):
        row["ordernum"] = orders
    if _plaintext_metric(page_views):
        row["impression"] = page_views
    return row


def _sydney_blocks(node: Any, found: list[dict], depth: int = 0) -> None:
    if depth > 8:
        return
    if isinstance(node, dict):
        if any(key in node for key in ("yesData", "todayData", "yesList", "todayList")):
            found.append(node)
            return
        for value in node.values():
            if isinstance(value, (dict, list)):
                _sydney_blocks(value, found, depth + 1)
    elif isinstance(node, list):
        for item in node[:20]:
            _sydney_blocks(item, found, depth + 1)


def _home_spend_value(item: dict) -> Any:
    for key, value in item.items():
        if _key_norm(key) in _HOME_SPEND_KEYS and _plaintext_metric(value):
            return value
    return None


def _home_spend_row(amount: Any, day: date, *, source: str) -> tuple[str, dict] | None:
    if not _plaintext_metric(amount):
        return None
    return (
        "goods",
        {
            "goodsid": "home-promo-spend",
            "goodsname": "首页推广花费",
            "statdate": day.isoformat(),
            "scene": "商品数据",
            "spend": amount,
            "__source": source,
        },
    )


def parse_home_promo_spend_text(text: str) -> tuple[str | None, str | None]:
    """首页可见文案。推广花费旁边的今日、昨日金额按元，不当成测款。"""
    compact = re.sub(r"\s+", " ", text or "")
    amount = r"((?:[0-9]{1,3}(?:,[0-9]{3})+|[0-9]+)(?:\.[0-9]{1,2})?)(?![\d.])(?!%)"
    today_amount = None
    yesterday_amount = None
    for match in re.finditer("推广花费", compact):
        window = compact[max(0, match.start() - 80) : match.end() + 180]
        today = re.search(r"(?:今日|今天)[^\d%]{0,16}" + amount, window)
        yesterday = re.search(r"(?:昨日|昨天)[^\d%]{0,16}" + amount, window)
        if today:
            today_amount = today.group(1).replace(",", "")
        if yesterday:
            yesterday_amount = yesterday.group(1).replace(",", "")
        if today_amount or yesterday_amount:
            return today_amount, yesterday_amount
    return None, None


def extract_home_spend_rows(data: Any, today: date) -> list[tuple[str, dict]]:
    """首页汇总里的推广花费。没有商品或创意身份，不能写成测款、测图。"""
    rows: list[tuple[str, dict]] = []
    yesterday = today - timedelta(days=1)
    seen: set[str] = set()

    def add(amount: Any, day: date, source: str) -> None:
        mapped = _home_spend_row(amount, day, source=source)
        if mapped is None:
            return
        stat = mapped[1]["statdate"]
        if stat in seen:
            return
        seen.add(stat)
        rows.append(mapped)

    def walk(node: Any, depth: int) -> None:
        if depth > 8:
            return
        if isinstance(node, dict):
            source = "ui" if node.get("source") == "mms-home-ui" else "json"
            if "todayPromoSpend" in node or "yesPromoSpend" in node:
                add(node.get("todayPromoSpend"), today, source)
                add(node.get("yesPromoSpend"), yesterday, source)
            for key, day in (("todayData", today), ("yesData", yesterday), ("today", today), ("yesterday", yesterday)):
                block = node.get(key)
                if isinstance(block, dict):
                    amount = _home_spend_value(block)
                    if amount is not None:
                        add(amount, day, "json")
            label = ""
            for key in ("title", "name", "cardName", "label", "indexName", "quotaName"):
                value = node.get(key)
                if isinstance(value, str):
                    label = value
                    break
            if label.strip() == "推广花费":
                today_raw = None
                yesterday_raw = None
                for key in ("todayValue", "today", "curValue", "value"):
                    if _plaintext_metric(node.get(key)):
                        today_raw = node.get(key)
                        break
                for key in ("yesterdayValue", "yesterday", "yesValue", "compareValue"):
                    if _plaintext_metric(node.get(key)):
                        yesterday_raw = node.get(key)
                        break
                add(today_raw, today, "json")
                add(yesterday_raw, yesterday, "json")
            for value in node.values():
                if isinstance(value, (dict, list)):
                    walk(value, depth + 1)
        elif isinstance(node, list):
            for item in node[:30]:
                walk(item, depth + 1)

    walk(data, 0)
    return rows


def extract_sydney_rows(data: Any, today: date) -> list[tuple[str, dict]]:
    """商家后台商品数据。yesData 是昨天的明文成交额，不是推广花费。"""
    blocks: list[dict] = []
    _sydney_blocks(data, blocks)
    rows: list[tuple[str, dict]] = []
    yesterday = today - timedelta(days=1)
    for block in blocks:
        list_rows: list[tuple[str, dict]] = []
        for key, day in (("yesList", yesterday), ("todayList", today)):
            items = block.get(key)
            if not isinstance(items, list):
                continue
            for item in items:
                if not isinstance(item, dict):
                    continue
                mapped = _sydney_item_row(item, day, aggregate=False)
                if mapped:
                    list_rows.append(("goods", mapped))
        if list_rows:
            rows.extend(list_rows)
            continue
        for key, day in (("yesData", yesterday), ("todayData", today)):
            item = block.get(key)
            if not isinstance(item, dict):
                continue
            mapped = _sydney_item_row(item, day, aggregate=True)
            if mapped:
                rows.append(("goods", mapped))
    return rows


def normalize_captures(
    captures: list[dict],
    *,
    shop_hint: str = "",
    today: date,
) -> SyncPayload:
    rows: list[tuple[str, dict[str, Any]]] = []
    mall = {"id": "", "name": ""}
    notes: list[str] = []
    auth_failed = False
    token_invalid = False
    promo_metrics = False
    pending_home: list[tuple[str, dict]] = []
    for item in captures:
        kind = item.get("kind")
        body = item.get("body")
        if kind == "note" and body:
            text = str(body).strip()[:300]
            notes.append(text)
            if note_says_auth_failed(text):
                auth_failed = True
            continue
        if kind == "debug":
            continue
        if kind == "json":
            data = _load_json(body)
            if data is None:
                continue
            if body_token_invalid(data):
                token_invalid = True
                auth_failed = True
            elif body_auth_failed(data):
                auth_failed = True
            if json_has_metrics(data) and not body_auth_failed(data):
                promo_metrics = True
            sydney_rows = extract_sydney_rows(data, today)
            if sydney_rows:
                rows.extend(sydney_rows)
                if SYDNEY_SOURCE_NOTE not in notes:
                    notes.append(SYDNEY_SOURCE_NOTE)
            pending_home.extend(extract_home_spend_rows(data, today))
            _walk(data, rows, mall, "json")
        elif kind == "csv":
            text = _decode_table(body)
            if text is None:
                notes.append("有一份 CSV 无法识别编码")
                continue
            for raw in _csv_rows(text):
                mapped = _map_flat(raw)
                _absorb_mall(mapped, mall)
                row_kind = _kind(mapped)
                if row_kind and _has_metric(mapped):
                    rows.append((row_kind, {**mapped, "__source": "csv"}))
                    promo_metrics = True
        elif kind == "xlsx":
            table, error = _xlsx_rows(body)
            if error:
                notes.append(error)
                continue
            for raw in table:
                mapped = _map_flat(raw)
                _absorb_mall(mapped, mall)
                row_kind = _kind(mapped)
                if row_kind and _has_metric(mapped):
                    rows.append((row_kind, {**mapped, "__source": "xlsx"}))
                    promo_metrics = True
    home_spend = False
    if pending_home and not promo_metrics:
        rows.extend(pending_home)
        home_spend = True
        if HOME_SPEND_NOTE not in notes:
            notes.append(HOME_SPEND_NOTE)
    if auth_failed and not promo_metrics and not any(note_says_auth_failed(item) for item in notes):
        notes.append(TOKEN_INVALID_NOTE if token_invalid else AUTH_FAILED_NOTE)
    if token_invalid and not promo_metrics:
        notes = [TOKEN_INVALID_NOTE if note == AUTH_FAILED_NOTE else note for note in notes]
        if not any("Token不合法" in item for item in notes):
            notes.insert(0, TOKEN_INVALID_NOTE)
    if home_spend:
        notes = [item.replace("这次没有写入推广花费。", "首页推广花费已另行写入，但测图测款仍未写入。") for item in notes]
    return _assemble(rows, mall, shop_hint=shop_hint, today=today, notes=notes)


def _load_json(body: Any) -> Any:
    if isinstance(body, (dict, list)):
        return body
    if isinstance(body, (bytes, bytearray)):
        try:
            body = body.decode("utf-8")
        except UnicodeDecodeError:
            return None
    if isinstance(body, str):
        try:
            return json.loads(body)
        except json.JSONDecodeError:
            return None
    return None


def _walk(node: Any, rows: list[tuple[str, dict]], mall: dict[str, str], source: str, depth: int = 0) -> None:
    if depth > 12:
        return
    if isinstance(node, list):
        for item in node:
            _walk(item, rows, mall, source, depth + 1)
        return
    if not isinstance(node, dict):
        return
    mapped = _map_flat(node)
    mapped["__source"] = source
    _absorb_mall(mapped, mall)
    nested = any(isinstance(value, (dict, list)) for value in node.values())
    row_kind = _kind(mapped)
    if row_kind and _has_metric(mapped) and (_scalar(mapped, _DATE_KEYS) not in (None, "") or not nested):
        rows.append((row_kind, mapped))
    for value in node.values():
        if isinstance(value, (dict, list)):
            _walk(value, rows, mall, source, depth + 1)


def _map_flat(raw: dict) -> dict[str, Any]:
    mapped: dict[str, Any] = {}
    for key, value in raw.items():
        if isinstance(value, (dict, list)):
            continue
        mapped[_norm(key)] = value
    return mapped


def _norm(key: Any) -> str:
    return str(key).strip().lower().replace("_", "").replace(" ", "")


def _scalar(row: dict[str, Any], keys: tuple[str, ...]) -> Any:
    for key in keys:
        value = row.get(_norm(key))
        if isinstance(value, (dict, list)) or value is None or value == "":
            continue
        return value
    return None


def _kind(row: dict[str, Any]) -> str | None:
    if _scalar(row, _CREATIVE_ID) is not None or _scalar(row, _CREATIVE_NAME) is not None:
        return "image"
    if _scalar(row, _GOODS_ID) is not None or _scalar(row, _GOODS_NAME) is not None:
        return "goods"
    if _scalar(row, _PLAN_ID) is not None or _scalar(row, _PLAN_NAME) is not None:
        return "campaign"
    return None


def _has_metric(row: dict[str, Any]) -> bool:
    for keys in (_IMPR, _CLICK, _SPEND, _ORDER, _GMV):
        if _scalar(row, keys) is not None:
            return True
    return False


def _absorb_mall(row: dict[str, Any], mall: dict[str, str]) -> None:
    if not mall["id"]:
        raw = _scalar(row, _MALL_ID)
        if raw is not None:
            mall["id"] = _id_token(raw)
    if not mall["name"]:
        raw = _scalar(row, _MALL_NAME)
        if raw is not None:
            mall["name"] = str(raw).strip()[:40]


def _decode_table(body: Any) -> str | None:
    if isinstance(body, str):
        return body
    if not isinstance(body, (bytes, bytearray)):
        return None
    for encoding in ("utf-8-sig", "utf-8", "gbk"):
        try:
            return bytes(body).decode(encoding)
        except UnicodeDecodeError:
            continue
    return None


def _csv_rows(text: str) -> list[dict[str, str]]:
    sample = text.lstrip("\ufeff")
    try:
        return list(csv.DictReader(io.StringIO(sample)))
    except csv.Error:
        return []


def _xlsx_rows(body: Any) -> tuple[list[dict], str]:
    if not isinstance(body, (bytes, bytearray)):
        return [], ""
    try:
        from openpyxl import load_workbook
    except ImportError:
        return [], "导出了 xlsx，但本机没有 openpyxl，这份表格被跳过"
    try:
        book = load_workbook(io.BytesIO(bytes(body)), read_only=True, data_only=True)
    except Exception:
        return [], "有一份 xlsx 打不开"
    try:
        sheet = book.active
        grid = list(sheet.iter_rows(values_only=True))
    finally:
        book.close()
    if not grid:
        return [], ""
    headers = ["" if cell is None else str(cell).strip() for cell in grid[0]]
    rows = []
    for line in grid[1:]:
        if line is None or all(cell is None or cell == "" for cell in line):
            continue
        rows.append({headers[index]: line[index] for index in range(min(len(headers), len(line))) if headers[index]})
    return rows, ""


def _assemble(
    rows: list[tuple[str, dict]],
    mall: dict[str, str],
    *,
    shop_hint: str,
    today: date,
    notes: list[str],
) -> SyncPayload:
    goods_meta: dict[str, dict] = {}
    goods_days: dict[tuple[str, str], MetricDay] = {}
    goods_scene: dict[str, str] = {}
    goods_plan: dict[str, str] = {}
    plan_meta: dict[str, dict] = {}
    plan_days: dict[tuple[str, str], MetricDay] = {}
    plan_by_name: dict[str, str] = {}
    image_meta: dict[str, dict] = {}
    image_days: dict[tuple[str, str], MetricDay] = {}
    explicit_plan = False

    for kind, row in rows:
        source = str(row.get("__source") or "json")
        stat, metric = _metric(row, source, today)
        if kind == "campaign":
            explicit_plan = True
            plan_id = _entity_id("live-plan-", _scalar(row, _PLAN_ID), _scalar(row, _PLAN_NAME) or "计划")
            name = str(_scalar(row, _PLAN_NAME) or plan_id)
            plan_meta[plan_id] = {
                "name": name[:80],
                "scene": str(_scalar(row, _SCENE) or "商品推广")[:40],
                "status": _status(_scalar(row, _STATUS)),
                "budget": money_to_fen(_scalar(row, _BUDGET), source=source),
                "bid": _optional_money(row, source),
            }
            plan_by_name[name] = plan_id
            plan_days[(plan_id, stat)] = metric
            continue
        if kind == "goods":
            goods_id = _entity_id("live-goods-", _scalar(row, _GOODS_ID), _scalar(row, _GOODS_NAME) or "商品")
            scene = str(_scalar(row, _SCENE) or "商品推广")[:40]
            goods_meta[goods_id] = {
                "name": str(_scalar(row, _GOODS_NAME) or goods_id)[:80],
                "scene": scene,
                "price": money_to_fen(_scalar(row, _PRICE), source=source),
                "stage": _stage(_scalar(row, _STAGE)),
                "budget": money_to_fen(_scalar(row, _BUDGET), source=source),
            }
            goods_scene[goods_id] = scene
            plan_raw = _scalar(row, _PLAN_ID)
            plan_name = _scalar(row, _PLAN_NAME)
            if plan_raw is not None:
                goods_plan[goods_id] = _entity_id("live-plan-", plan_raw, str(plan_name or plan_raw))
            elif plan_name is not None:
                goods_plan[goods_id] = str(plan_name)
            goods_days[(goods_id, stat)] = metric
            continue
        image_id = _entity_id("live-img-", _scalar(row, _CREATIVE_ID), _scalar(row, _CREATIVE_NAME) or "主图")
        goods_raw = _scalar(row, _GOODS_ID)
        goods_name = _scalar(row, _GOODS_NAME)
        linked = _entity_id("live-goods-", goods_raw, str(goods_name or "未命名商品")) if goods_raw is not None or goods_name is not None else "live-goods-unknown"
        image_meta[image_id] = {
            "label": str(_scalar(row, _CREATIVE_NAME) or "主图")[:40],
            "goods_id": linked,
            "goods_name": str(goods_name or "")[:80],
        }
        image_days[(image_id, stat)] = metric

    for image in image_meta.values():
        gid = image["goods_id"]
        if gid not in goods_meta:
            goods_meta[gid] = {
                "name": image["goods_name"] or "未命名商品",
                "scene": "商品推广",
                "price": 0,
                "stage": "推广中",
                "budget": 0,
            }
            goods_scene[gid] = "商品推广"

    if not explicit_plan:
        _synthesize_plans(goods_meta, goods_days, goods_scene, plan_meta, plan_days)
        for gid in goods_meta:
            goods_plan[gid] = f"live-scene-{_slug(goods_scene.get(gid) or '商品推广')}"
    else:
        fallback = next(iter(plan_meta), "")
        for gid, meta in goods_meta.items():
            linked = goods_plan.get(gid, "")
            if linked in plan_meta:
                continue
            if linked and linked in plan_by_name:
                goods_plan[gid] = plan_by_name[linked]
                continue
            if fallback:
                goods_plan[gid] = fallback
                continue
            scene = goods_scene.get(gid) or "商品推广"
            plan_id = f"live-scene-{_slug(scene)}"
            plan_meta.setdefault(
                plan_id,
                {"name": scene, "scene": scene, "status": "投放中", "budget": 0, "bid": None},
            )
            goods_plan[gid] = plan_id

    mall_id = mall["id"]
    shop_id = f"live-{_slug(mall_id)}" if mall_id else "live-unknown"
    name = (mall["name"] or shop_hint or "拼多多店铺")[:40]
    note = " ".join(notes)
    if not goods_meta and not plan_meta and not note_says_auth_failed(note) and "登录已保存" not in note:
        extra = (
            "已登录，但没有从报表里识别到推广商品或计划。"
            + manual_login_tail("推广中心页面")
            + "测图只有报表带创意或主图时才会出现。"
        )
        note = f"{note} {extra}".strip()

    campaigns = [
        CampaignInfo(
            campaign_id=plan_id,
            name=meta["name"],
            scene=meta["scene"],
            status=meta["status"],
            daily_budget_fen=meta["budget"],
            bid_fen=meta["bid"],
            pace="normal",
        )
        for plan_id, meta in plan_meta.items()
    ]
    campaign_days = _group_days(plan_days)
    goods = [
        GoodsInfo(
            goods_id=gid,
            campaign_id=goods_plan[gid],
            name=meta["name"],
            category=meta["scene"],
            price_fen=meta["price"],
            stage=meta["stage"],
            daily_budget_fen=meta["budget"],
        )
        for gid, meta in goods_meta.items()
    ]
    images = [
        ImageInfo(image_id=image_id, goods_id=meta["goods_id"], label=meta["label"], tone="plain")
        for image_id, meta in image_meta.items()
    ]
    return SyncPayload(
        shop=ShopProfile(shop_id=shop_id, mall_id=mall_id, name=name),
        campaigns=campaigns,
        campaign_days=campaign_days,
        goods=goods,
        goods_days=_group_days(goods_days),
        images=images,
        image_days=_group_days(image_days),
        note=note,
    )


def _synthesize_plans(goods_meta, goods_days, goods_scene, plan_meta, plan_days) -> None:
    scenes = {goods_scene.get(gid) or "商品推广" for gid in goods_meta} or set()
    if not scenes and not goods_meta:
        return
    buckets: dict[tuple[str, str], list[MetricDay]] = {}
    budgets: dict[str, int] = {}
    for gid, meta in goods_meta.items():
        scene = goods_scene.get(gid) or "商品推广"
        budgets[scene] = budgets.get(scene, 0) + int(meta["budget"])
        for (owner, stat), metric in goods_days.items():
            if owner != gid:
                continue
            buckets.setdefault((scene, stat), []).append(metric)
    for scene in scenes:
        plan_id = f"live-scene-{_slug(scene)}"
        plan_meta[plan_id] = {
            "name": scene,
            "scene": scene,
            "status": "投放中",
            "budget": budgets.get(scene, 0),
            "bid": None,
        }
        for (bucket_scene, stat), metrics in buckets.items():
            if bucket_scene != scene:
                continue
            plan_days[(plan_id, stat)] = _sum_metrics(stat, metrics)


def _metric(row: dict[str, Any], source: str, today: date) -> tuple[str, MetricDay]:
    stat = _date_text(_scalar(row, _DATE_KEYS)) or today.isoformat()
    return stat, MetricDay(
        stat_date=stat,
        impressions=_count(_scalar(row, _IMPR)),
        clicks=_count(_scalar(row, _CLICK)),
        spend_fen=money_to_fen(_scalar(row, _SPEND), source=source),
        orders=_count(_scalar(row, _ORDER)),
        gmv_fen=money_to_fen(_scalar(row, _GMV), source=source),
        add_to_cart=_count(_scalar(row, _CART)),
    )


def _sum_metrics(stat: str, metrics: list[MetricDay]) -> MetricDay:
    return MetricDay(
        stat_date=stat,
        impressions=sum(item.impressions for item in metrics),
        clicks=sum(item.clicks for item in metrics),
        spend_fen=sum(item.spend_fen for item in metrics),
        orders=sum(item.orders for item in metrics),
        gmv_fen=sum(item.gmv_fen for item in metrics),
        add_to_cart=sum(item.add_to_cart for item in metrics),
    )


def _group_days(store: dict[tuple[str, str], MetricDay]) -> dict[str, list[MetricDay]]:
    grouped: dict[str, list[MetricDay]] = {}
    for (owner, _stat), metric in store.items():
        grouped.setdefault(owner, []).append(metric)
    for series in grouped.values():
        series.sort(key=lambda item: item.stat_date)
    return grouped


def _optional_money(row: dict[str, Any], source: str) -> int | None:
    raw = _scalar(row, _BID)
    if raw is None:
        return None
    return money_to_fen(raw, source=source)


def _status(value: Any) -> str:
    text = str(value or "").strip()
    lowered = text.lower()
    if lowered in {"1", "running", "online", "投放中", "推广中", "进行中", "启用"}:
        return "投放中"
    if lowered in {"2", "paused", "pause", "暂停", "已暂停", "停止", "下线"}:
        return "已暂停"
    if not text:
        return "投放中"
    return text[:20]


def _stage(value: Any) -> str:
    text = str(value or "").strip()
    if not text:
        return "推广中"
    if "积累" in text:
        return "数据积累期"
    if "智能" in text or "学习" in text:
        return "学习期"
    return text[:20]


def _count(value: Any) -> int:
    if value is None or isinstance(value, bool) or value == "":
        return 0
    if isinstance(value, str):
        value = value.replace(",", "").strip()
        if value in {"-", "--"}:
            return 0
    try:
        return max(0, int(float(value)))
    except (TypeError, ValueError):
        return 0


def _date_text(value: Any) -> str | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = str(value).strip().replace("/", "-").replace(".", "-")
    if " " in text:
        text = text.split(" ", 1)[0]
    if "T" in text:
        text = text.split("T", 1)[0]
    if re.fullmatch(r"\d{8}", text):
        return f"{text[:4]}-{text[4:6]}-{text[6:]}"
    match = re.fullmatch(r"(\d{4})-(\d{1,2})-(\d{1,2})", text)
    if match:
        year, month, day = match.groups()
        return f"{int(year):04d}-{int(month):02d}-{int(day):02d}"
    return None


def _id_token(value: Any) -> str:
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, int) and not isinstance(value, bool):
        return str(value)
    return str(value).strip()


def _entity_id(prefix: str, raw: Any, fallback: str) -> str:
    token = _slug(_id_token(raw)) if raw not in (None, "") else _slug(str(fallback))
    return f"{prefix}{token}"


def _slug(text: str) -> str:
    cleaned = re.sub(r"\s+", "-", text.strip())
    cleaned = re.sub(r"[^0-9A-Za-z\u4e00-\u9fff_-]", "", cleaned)
    return cleaned[:40] or "item"
